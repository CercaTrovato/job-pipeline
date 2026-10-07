from __future__ import annotations

import io
import json
import pathlib
import re
import uuid
import zipfile
from typing import Any, Dict, List

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

from jp import packets

MAX_UPLOAD = 5 * 1024 * 1024
MAX_UNPACKED = 20 * 1024 * 1024
_SCHEMA = pathlib.Path(__file__).parent / "resources" / "resume.schema.json"
_SPEC = pathlib.Path(__file__).resolve().parents[1] / "llm" / "specs" / "resume" / "SKILL.md"
_PACKAGED_SPEC = pathlib.Path(__file__).parent / "resources" / "llm" / "specs" / "resume" / "SKILL.md"


class ResumeError(ValueError):
    """简历输入或草稿不符合安全、来源和结构约束。"""


def extract_text(data: bytes, filename: str) -> str:
    """只在内存解析 PDF/DOCX；不写入上传文件或解包文件。"""
    if not isinstance(data, bytes) or len(data) > MAX_UPLOAD:
        raise ResumeError("文件超过 5 MB 或输入格式无效")
    suffix = pathlib.Path(filename or "").suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ResumeError("不支持加密 PDF")
            if not reader.pages or len(reader.pages) > 20:
                raise ResumeError("PDF 页数必须在 1 到 20 页之间")
            pages = []
            total_bytes = 0
            for page in reader.pages:
                page_text = (page.extract_text() or "").strip()
                total_bytes += len(page_text.encode("utf-8"))
                if total_bytes > MAX_UNPACKED:
                    raise ResumeError("解包文本超过 20 MB")
                pages.append(page_text)
            if any(len(page) < 8 for page in pages):
                raise ResumeError("PDF 含有无可提取文本的页面，可能是扫描件；请提供可选中文本版")
            text = "\n\n".join(pages)
        elif suffix == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_UNPACKED:
                    raise ResumeError("DOCX 解包大小超过 20 MB")
                if "word/document.xml" not in archive.namelist():
                    raise ResumeError("DOCX 文件结构不完整")
            doc = Document(io.BytesIO(data))
            body = doc.element.body
            explicit_breaks = sum(1 for node in body.iter(qn("w:br"))
                                  if node.get(qn("w:type")) == "page")
            rendered_breaks = sum(1 for _ in body.iter(qn("w:lastRenderedPageBreak")))
            page_floor = 1 + max(explicit_breaks, rendered_breaks, len(doc.sections) - 1)
            if page_floor > 20:
                raise ResumeError("DOCX 页数超过 20 页")
            blocks = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append(" | ".join(cells))
            text = "\n".join(blocks).strip()
        else:
            raise ResumeError("仅支持 PDF 或 DOCX")
    except ResumeError:
        raise
    except Exception as exc:
        raise ResumeError("文件损坏或无法解析") from exc
    if len(text.encode("utf-8")) > MAX_UNPACKED:
        raise ResumeError("解包文本超过 20 MB")
    if len(re.sub(r"\s", "", text)) < 20:
        raise ResumeError("文件没有足够的可提取文本，可能是扫描件或空文件")
    return text


_PII_PATTERNS = (
    ("email", re.compile(r"(?<![\w.+-])[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}(?![\w.-])", re.I)),
    ("id", re.compile(r"(?<!\d)(?:\d{17}[\dXx]|\d{15})(?!\d)")),
    ("phone", re.compile(r"(?<!\d)(?:\+?\d{1,3}[\s().-]?)?(?:1[3-9]\d{9}|\d{3}[\s.-]\d{3}[\s.-]\d{4})(?!\d)")),
    ("age", re.compile(r"(?i)(?:(?:年龄|年齡|age)\s*[:：]?\s*\d{1,3}\s*(?:岁|歲|years?\s*old)?|\b\d{1,3}\s*(?:岁|歲|years?\s*old)\b)")),
    ("gender", re.compile(r"(?i)(?:(?:性别|性別|gender|sex)\s*[:：]?\s*(?:男性|女性|男|女|male|female|non[- ]?binary|other)|\b(?:male|female|non[- ]?binary)\b)")),
)


def redact(text: str) -> str:
    """确定性遮盖常见联系方式和年龄、性别字段，供用户检查本地预览。"""
    if not isinstance(text, str):
        raise ResumeError("简历文本必须是字符串")
    result = text
    for label, pattern in _PII_PATTERNS:
        result = pattern.sub("[已脱敏:%s]" % label, result)
    return result


def _atomic_json(path: pathlib.Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp-" + uuid.uuid4().hex)
    try:
        with tmp.open("w", encoding="utf-8", newline="\n") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
        tmp.replace(path)
    finally:
        if tmp.exists():
            tmp.unlink()


def _atomic_text(path: pathlib.Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp-" + uuid.uuid4().hex)
    try:
        tmp.write_text(value, encoding="utf-8", newline="\n")
        tmp.replace(path)
    finally:
        if tmp.exists():
            tmp.unlink()


def _read_json(path: pathlib.Path) -> Any:
    with path.open("r", encoding="utf-8") as stream:
        return json.load(stream)


def _keywords(text: str) -> List[str]:
    try:
        from jp.facts_index import tokenize
        return tokenize(text)
    except ImportError:
        words = re.findall(r"[A-Za-z][A-Za-z0-9+#._-]{1,}|[\u4e00-\u9fff]{2,}", text)
        return list(dict.fromkeys(word.lower() for word in words))


def _workspace_path(workspace: Any) -> pathlib.Path:
    if not isinstance(workspace, (str, pathlib.Path)) and hasattr(workspace, "root"):
        workspace = workspace.root
    if isinstance(workspace, dict):
        workspace = workspace.get("workspace") or workspace.get("paths", {}).get("workspace")
    if not workspace:
        raise ResumeError("需要传入用户 workspace 路径")
    return pathlib.Path(workspace).expanduser().resolve()


def _facts_index_path(workspace: Any, root: pathlib.Path) -> pathlib.Path:
    cfg = None
    if hasattr(workspace, "runtime_config"):
        cfg = workspace.runtime_config()
    elif isinstance(workspace, dict):
        cfg = workspace
    configured = ((cfg or {}).get("paths") or {}).get("facts_index")
    if configured:
        path = pathlib.Path(configured).expanduser()
        return path if path.is_absolute() else (root / path).resolve()
    return root / "data" / "facts_index.json"


def read_draft(workspace) -> dict:
    return _read_json(_workspace_path(workspace) / "profile" / "resume-draft.json")


def _validate_items(items: Any, source_text: str) -> List[Dict[str, Any]]:
    if not isinstance(items, list):
        raise ResumeError("草稿 items 必须是数组")
    ids = set()
    clean = []
    for raw in items:
        if not isinstance(raw, dict):
            raise ResumeError("草稿条目格式无效")
        item = dict(raw)
        if item.get("id") in ids:
            raise ResumeError("草稿条目 ID 重复")
        if not isinstance(item.get("id"), str) or redact(item["id"]) != item["id"]:
            raise ResumeError("草稿条目 ID 含有待脱敏信息")
        ids.add(item.get("id"))
        if item.get("source_quote") not in source_text:
            raise ResumeError("source_quote 不在脱敏原文中")
        if item.get("confirmed") is not False:
            raise ResumeError("模型草稿不得自行确认条目")
        for field in ("title", "text", "source_quote"):
            val = item.get(field)
            if not isinstance(val, str) or redact(val) != val:
                raise ResumeError("草稿含有待脱敏信息，请重新检查")
        clean.append(item)
    return clean


def make_draft(workspace, text: str, cfg: dict, generate: bool = False) -> dict:
    """保存本地预览；可选地为已脱敏且經用户校正的文本创建模型任务包。"""
    root = _workspace_path(workspace)
    safe_text = redact(text)
    if generate and safe_text != text:
        raise ResumeError("请先检查并校正脱敏预览，再将校正后的脱敏文本传入 generate=True")
    draft_path = root / "profile" / "resume-draft.json"
    existing = _read_json(draft_path) if draft_path.exists() else {}
    same_text = existing.get("text") == safe_text
    draft = {"version": 1, "id": existing.get("id") if same_text and existing.get("id") else uuid.uuid4().hex,
             "text": safe_text, "items": existing.get("items", []) if same_text else [], "status": "preview",
             "privacy_notice": "预览仅遮盖常见联系方式、证件号、年龄和性别字段；姓名等其他敏感信息可能保留，请逐项检查并手工删除。"}
    _atomic_json(draft_path, draft)
    if not generate:
        return draft

    task_dir = root / "tasks" / "_resume" / "resume"
    task_dir.mkdir(parents=True, exist_ok=True)
    input_obj = {"text": safe_text}
    old_input = task_dir / "input.json"
    if old_input.exists() and packets.input_hash(_read_json(old_input)) != packets.input_hash(input_obj):
        for stale_name in ("output.json", "error.txt", "error.txt.json", "usage.json"):
            stale = task_dir / stale_name
            if stale.exists():
                stale.unlink()
    _atomic_json(task_dir / "input.json", input_obj)
    spec_path = _SPEC if _SPEC.exists() else _PACKAGED_SPEC
    if not spec_path.exists():
        raise ResumeError("简历任务说明资源不可用")
    prompt = spec_path.read_text(encoding="utf-8") + "\n\n## 输入文本\n\n" + safe_text
    _atomic_text(task_dir / "prompt.md", prompt)
    _atomic_text(task_dir / "schema.json", _SCHEMA.read_text(encoding="utf-8"))
    p = packets.Packet("_resume", "resume", task_dir, packets.input_hash(input_obj))
    if (task_dir / "output.json").exists():
        return collect_draft(workspace)
    backend = (cfg.get("llm") or {}).get("backend", "manual")
    if backend == "api":
        from jp import backends
        result = backends.run_api([p], backends.load_api_config(cfg))
        if result.failed:
            _atomic_text(task_dir / "error.txt", "api: 模型生成失败；请检查配置后重试。")
            raise RuntimeError("api 模型任务失败；请检查配置后重试。")
    elif backend == "codex":
        from jp import backends
        llm = cfg.get("llm") or {}
        model = llm.get("codex_model") or llm.get("model")
        if not model:
            raise RuntimeError("codex 后端必须配置模型")
        result = backends.run_codex([p], timeout=llm.get("codex_timeout_sec", llm.get("timeout_sec", 180)),
                                    model=model,
                                    reasoning_effort=llm.get("codex_reasoning_effort") or llm.get("reasoning_effort", "low"),
                                    concurrency=1)
        if result.failed:
            _atomic_text(task_dir / "error.txt", "codex: 模型生成失败；请检查配置后重试。")
            raise RuntimeError("codex 模型任务失败；请检查配置后重试。")
    elif backend not in ("manual", "claude"):
        raise ResumeError("不支持的模型后端")
    draft["status"] = "pending"
    draft["task_dir"] = str(task_dir)
    if backend in ("manual", "claude"):
        draft["manual"] = True
    _atomic_json(draft_path, draft)
    if backend in ("api", "codex"):
        if not (task_dir / "output.json").exists():
            raise RuntimeError("模型任务未生成有效输出。")
        return collect_draft(root)
    return draft


def collect_draft(workspace) -> dict:
    root = _workspace_path(workspace)
    draft_path = root / "profile" / "resume-draft.json"
    task_dir = root / "tasks" / "_resume" / "resume"
    draft = _read_json(draft_path)
    out_path = task_dir / "output.json"
    if not out_path.exists():
        raise ResumeError("简历草稿任务尚无 output.json")
    try:
        output = _read_json(out_path)
        import jsonschema
        jsonschema.validate(output, _read_json(task_dir / "schema.json"))
        items = _validate_items(output["items"], draft["text"])
    except ResumeError:
        raise
    except Exception as exc:
        # 不把模型输出或异常正文写入对外可见错误文件。
        _atomic_json(task_dir / "error.txt.json", {"error": "简历草稿校验失败，请重新生成或手工修正任务输出。"})
        raise ResumeError("简历草稿结构或内容校验失败") from exc
    draft["items"] = items
    draft["status"] = "ready_for_review"
    _atomic_json(draft_path, draft)
    return draft


def confirm_draft(workspace, draft_id: str, items: list) -> list:
    """保留用户确认的非推断事实，并写入带备份的本地 facts 索引。"""
    root = _workspace_path(workspace)
    draft_path = root / "profile" / "resume-draft.json"
    draft = _read_json(draft_path)
    if draft.get("id") != draft_id:
        raise ResumeError("简历草稿 ID 不匹配")
    original = {entry["id"]: entry for entry in draft.get("items", [])}
    if not isinstance(items, list):
        raise ResumeError("确认条目必须是数组")
    accepted = []
    seen = set()
    for submitted in items:
        if not isinstance(submitted, dict) or submitted.get("id") not in original:
            raise ResumeError("确认条目不属于当前草稿")
        id_ = submitted["id"]
        if id_ in seen:
            raise ResumeError("确认条目 ID 重复")
        seen.add(id_)
        source = original[id_]
        if source.get("inferred") or not submitted.get("confirmed"):
            continue
        quote = source.get("source_quote", "")
        title, fact_text = submitted.get("title"), submitted.get("text")
        if not quote or quote not in draft.get("text", ""):
            raise ResumeError("stored source_quote 不在当前脱敏原文中")
        if not isinstance(title, str) or not isinstance(fact_text, str):
            raise ResumeError("确认事实的标题或文本无效")
        title, fact_text = redact(title).strip(), redact(fact_text).strip()
        if not title or not fact_text:
            raise ResumeError("确认事实的标题或文本不能为空")
        accepted.append({"id": id_, "type": source["type"], "title": title, "text": fact_text,
                         "keywords": _keywords(title + " " + fact_text), "grade": "C",
                         "source_quote": redact(quote)})

    facts_path = _facts_index_path(workspace, root)
    previous = _read_json(facts_path) if facts_path.exists() else []
    backup = facts_path.with_name("facts-index.backup-%s.json" % uuid.uuid4().hex[:12])
    if facts_path.exists():
        backup.write_bytes(facts_path.read_bytes())
    new_by_id = {entry["id"]: entry for entry in accepted}
    merged = [new_by_id.pop(entry.get("id"), entry) for entry in previous if isinstance(entry, dict)]
    merged.extend(new_by_id.values())
    _atomic_json(facts_path, merged)
    return accepted

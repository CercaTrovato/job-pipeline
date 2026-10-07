from __future__ import annotations

import io
import json
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from pypdf import PdfWriter

from jp import resume


def _docx_bytes():
    doc = Document()
    doc.add_paragraph("Python 数据分析项目经验")
    doc.add_paragraph("熟悉 SQL，完成课程研究项目")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pdf_bytes(encrypted=False):
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    if encrypted:
        writer.encrypt("password")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def test_extract_docx_in_memory_and_reject_scanned_pdf():
    text = resume.extract_text(_docx_bytes(), "resume.docx")
    assert "Python 数据分析项目经验" in text
    with pytest.raises(resume.ResumeError, match="扫描件"):
        resume.extract_text(_pdf_bytes(), "scan.pdf")


def test_extract_rejects_encrypted_oversized_and_unsupported():
    with pytest.raises(resume.ResumeError, match="加密"):
        resume.extract_text(_pdf_bytes(encrypted=True), "resume.pdf")
    with pytest.raises(resume.ResumeError, match="5 MB"):
        resume.extract_text(b"x" * (5 * 1024 * 1024 + 1), "resume.pdf")
    with pytest.raises(resume.ResumeError, match="仅支持"):
        resume.extract_text(b"data", "resume.doc")


def test_extract_rejects_docx_over_twenty_explicit_pages():
    doc = Document()
    for _ in range(20):
        doc.add_paragraph("page").add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("last page")
    out = io.BytesIO()
    doc.save(out)
    with pytest.raises(resume.ResumeError, match="20 页"):
        resume.extract_text(out.getvalue(), "many-pages.docx")


def test_redact_sensitive_fields_and_keep_manual_preview(tmp_path):
    source = "电话 13800138000 邮箱 a@example.com 身份证 11010519491231002X 年龄：28 性别：女"
    safe = resume.redact(source)
    assert all(value not in safe for value in ("13800138000", "a@example.com", "11010519491231002X", "28", "女"))
    draft = resume.make_draft(tmp_path, source, {"llm": {"backend": "manual"}})
    assert draft["status"] == "preview" and draft["text"] == safe
    assert json.loads((tmp_path / "profile" / "resume-draft.json").read_text(encoding="utf-8")) == draft
    with pytest.raises(resume.ResumeError, match="校正"):
        resume.make_draft(tmp_path, source, {"llm": {"backend": "manual"}}, generate=True)


def test_manual_task_collect_confirm_and_keep_backup(tmp_path):
    safe = "熟悉 Python，完成课程研究项目"
    draft = resume.make_draft(tmp_path, safe, {"llm": {"backend": "manual"}}, generate=True)
    packet_dir = tmp_path / "tasks" / "_resume" / "resume"
    output = {"items": [{"id": "skill-python", "type": "skill", "title": "Python", "text": "熟悉 Python",
                         "source_quote": "熟悉 Python", "confirmed": False, "inferred": False}]}
    (packet_dir / "output.json").write_text(json.dumps(output, ensure_ascii=False), encoding="utf-8")
    ready = resume.collect_draft(tmp_path)
    assert ready["status"] == "ready_for_review"
    accepted = resume.confirm_draft(tmp_path, draft["id"], [{**output["items"][0], "confirmed": True}])
    assert accepted[0]["grade"] == "C" and accepted[0]["source_quote"] == "熟悉 Python"
    facts = json.loads((tmp_path / "data" / "facts_index.json").read_text(encoding="utf-8"))
    assert facts[0]["id"] == "skill-python" and facts[0]["keywords"]
    resume.confirm_draft(tmp_path, draft["id"], [{**output["items"][0], "confirmed": True}])
    assert list((tmp_path / "data").glob("facts-index.backup-*.json"))


def test_collect_rejects_unanchored_quote_and_duplicate_ids(tmp_path):
    draft = resume.make_draft(tmp_path, "熟悉 Python 和 SQL", {"llm": {"backend": "manual"}}, generate=True)
    packet_dir = tmp_path / "tasks" / "_resume" / "resume"
    base = {"id": "same", "type": "skill", "title": "技能", "text": "熟悉 Python",
            "source_quote": "熟悉 Python", "confirmed": False, "inferred": False}
    (packet_dir / "output.json").write_text(json.dumps({"items": [base, base]}, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(resume.ResumeError, match="重复"):
        resume.collect_draft(tmp_path)
    base["source_quote"] = "不存在的引文"
    (packet_dir / "output.json").write_text(json.dumps({"items": [base]}, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(resume.ResumeError, match="source_quote 不在脱敏原文中"):
        resume.collect_draft(tmp_path)
